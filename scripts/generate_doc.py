import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_guide_docx(filepath):
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # --- Title Banner ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_sub = title_p.add_run("ESME SUDRIA • CYCLE INGÉNIEUR SPÉ (2026-2027)\n")
    run_sub.font.size = Pt(10)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5) # Indigo

    run_title = title_p.add_run("GUIDE D'UTILISATION ET DE TRANSMISSION DU PORTAIL")
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Dark Navy

    desc_p = doc.add_paragraph("Ce document récapitule tout ce qu'il faut savoir pour ouvrir et utiliser le portail web des cours, que vous soyez sous Windows ou sous Mac.")
    desc_p.paragraph_format.space_after = Pt(16)
    desc_p.runs[0].font.italic = True
    desc_p.runs[0].font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    # --- Section 1: Comment le transmettre à un ami ---
    h1 = doc.add_heading("1. Comment envoyer le portail à un ami ?", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.color.rgb = RGBColor(0x1E, 0x1B, 0x4B)
        r.font.size = Pt(14)

    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(8)
    p1.add_run("Le portail est ").font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    p1.add_run("100% autonome et portable").bold = True
    p1.add_run(" : il ne nécessite aucune connexion à internet pour lire les documents, et ne dépend d'aucun chemin d'accès figé.\n")
    
    steps = [
        ("Faire un clic droit", "sur le dossier « Site 2026-2027 »."),
        ("Choisir « Compresser dans un fichier ZIP »", "(ou « Envoyer vers > Dossier compressé »)."),
        ("Transmettre le fichier ZIP", "à votre ami par Google Drive, WeTransfer, Discord, Telegram, clé USB, etc."),
        ("Votre ami extrait le ZIP", "sur son propre ordinateur (Windows ou Mac) et suit les consignes ci-dessous.")
    ]
    for bold_text, rest in steps:
        sp = doc.add_paragraph(style='List Number')
        sp.paragraph_format.space_after = Pt(3)
        r1 = sp.add_run(bold_text + " ")
        r1.bold = True
        r1.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)
        sp.add_run(rest)

    # --- Section 2: Instructions sous WINDOWS ---
    h2 = doc.add_heading("2. Lancement sous WINDOWS (PC)", level=1)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    for r in h2.runs:
        r.font.color.rgb = RGBColor(0x1E, 0x1B, 0x4B)
        r.font.size = Pt(14)

    wp = doc.add_paragraph()
    wp.paragraph_format.space_after = Pt(6)
    r_w1 = wp.add_run("Méthode 1 : En un double-clic (Recommandée)\n")
    r_w1.bold = True
    r_w1.font.color.rgb = RGBColor(0x05, 0x96, 0x69) # Emerald
    wp.add_run("1. Dans le dossier « Site 2026-2027 », faites un simple double-clic sur ")
    wp.add_run("Lancer_Portail.bat").bold = True
    wp.add_run(".\n2. Une fenêtre noire démarre le serveur local et ouvre instantanément votre navigateur (Chrome ou Edge) à l'adresse : ")
    r_url = wp.add_run("http://localhost:8080\n")
    r_url.bold = True
    wp.add_run("3. Laissez la petite fenêtre ouverte en arrière-plan pendant que vous travaillez. Pour fermer le site quand vous avez terminé, fermez simplement cette fenêtre.\n")

    wp2 = doc.add_paragraph()
    wp2.paragraph_format.space_after = Pt(8)
    r_w2 = wp2.add_run("Méthode 2 : Sans aucun serveur (Ouverture directe)\n")
    r_w2.bold = True
    wp2.add_run("Vous pouvez également double-cliquer directement sur le fichier ")
    wp2.add_run("index.html").bold = True
    wp2.add_run(". Le portail s'ouvrira immédiatement dans votre navigateur habituel.")

    # --- Section 3: Instructions sous MAC / macOS ---
    h3 = doc.add_heading("3. Lancement sous MAC (Apple macOS) ou LINUX", level=1)
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    for r in h3.runs:
        r.font.color.rgb = RGBColor(0x1E, 0x1B, 0x4B)
        r.font.size = Pt(14)

    mp = doc.add_paragraph()
    mp.paragraph_format.space_after = Pt(6)
    r_m1 = mp.add_run("Méthode 1 : En un double-clic sur Mac\n")
    r_m1.bold = True
    r_m1.font.color.rgb = RGBColor(0x02, 0x84, 0xC7) # Sky
    mp.add_run("1. Dans le dossier, faites un double-clic sur le fichier ")
    mp.add_run("Lancer_Portail_Mac.command").bold = True
    mp.add_run(".\n2. Le terminal Mac lance le serveur et ouvre Safari ou Chrome à l'adresse ")
    mp.add_run("http://localhost:8080").bold = True
    mp.add_run(".\n*(Note : la première fois, si Mac bloque l'exécution, faites un clic droit sur le fichier > « Ouvrir » > « Ouvrir quand même »)*.\n")

    mp2 = doc.add_paragraph()
    mp2.paragraph_format.space_after = Pt(8)
    r_m2 = mp2.add_run("Méthode 2 : Directement dans le navigateur\n")
    r_m2.bold = True
    mp2.add_run("Faites un clic droit sur ")
    mp2.add_run("index.html").bold = True
    mp2.add_run(" > « Ouvrir avec » > « Safari » ou « Google Chrome ».")

    # --- Section 4: Ce qui est inclus dans le portail ---
    h4 = doc.add_heading("4. Contenu et Fonctionnalités du Portail", level=1)
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    for r in h4.runs:
        r.font.color.rgb = RGBColor(0x1E, 0x1B, 0x4B)
        r.font.size = Pt(14)

    # Summary table
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    headers = ["Unité d'Enseignement", "Matières Incluses", "Ressources Clés"]
    widths = [Inches(1.8), Inches(2.8), Inches(2.0)]
    
    for i, name in enumerate(headers):
        hdr_cells[i].text = name
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "4F46E5")
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.runs[0].font.size = Pt(10)

    rows_data = [
        ("UE1 : Mathématiques & Signal", "Maths Fondamentales 3, Outils Maths 3, Analyse Numérique, Signaux & Systèmes", "Cours, TD, Corrigés, Midterms"),
        ("UE2 : Sciences de l'Ingénieur", "Systèmes Technologiques 3, Mécanique des Fluides, Gestion de Projet, Enjeux Environnementaux 1", "Annales Spé, Corrigés Méca Flu, Templates"),
        ("UE3 : Numérique & Langues", "Algorithmique Avancée 2, Électronique Analogique 1, Anglais Professionnel", "Notebooks Python, Midterms EA1 Sujets + Corrigés"),
        ("UE4 : Entreprise & Économie", "Gestion d'Entreprise 1", "Comptabilité, Bilan, Finance")
    ]

    for ue, mat, res in rows_data:
        row_cells = table.add_row().cells
        row_cells[0].text = ue
        row_cells[0].width = widths[0]
        row_cells[1].text = mat
        row_cells[1].width = widths[1]
        row_cells[2].text = res
        row_cells[2].width = widths[2]
        
        for c in row_cells:
            p = c.paragraphs[0]
            p.runs[0].font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_before = Pt(8)

    features = [
        ("Distinction Paris vs National : ", "Bouton en haut de page pour n'afficher que les documents de Paris ou le tronc commun."),
        ("Onglet Examens & Midterms : ", "Accès instantané à tous les sujets de DS passés et à leurs corrections détaillées."),
        ("Visionneuse PDF intégrée : ", "Un clic sur « Consulter » ouvre le document dans le site avec les outils de lecture, zoom et impression."),
        ("Moteur de recherche instantané : ", "Tapez n'importe quel mot (ex: Fourier, Laplace, DS, Séries) ou utilisez le raccourci clavier /."),
        ("Mode Sombre / Mode Clair : ", "Bouton en haut à droite avec mémorisation de vos préférences.")
    ]

    for f_title, f_desc in features:
        fp = doc.add_paragraph(style='List Bullet')
        fp.paragraph_format.space_after = Pt(3)
        r = fp.add_run(f_title)
        r.bold = True
        r.font.color.rgb = RGBColor(0x31, 0x2E, 0x81)
        fp.add_run(f_desc)

    # --- Section 5: Mises à jour ---
    h5 = doc.add_heading("5. Comment ajouter de nouveaux cours plus tard ?", level=1)
    h5.paragraph_format.space_before = Pt(16)
    h5.paragraph_format.space_after = Pt(6)
    for r in h5.runs:
        r.font.color.rgb = RGBColor(0x1E, 0x1B, 0x4B)
        r.font.size = Pt(14)

    up = doc.add_paragraph()
    up.add_run("Si vous téléchargez de nouveaux PDF au cours du semestre, déposez-les simplement dans le dossier ")
    up.add_run("courses_data").bold = True
    up.add_run(" et double-cliquez sur ")
    up.add_run("Mettre_A_Jour.bat").bold = True
    up.add_run(". Le portail sera automatiquement réindexé et actualisé !\n")

    # Save
    doc.save(filepath)
    print(f"Document Word cree avec succes : {filepath}")

if __name__ == "__main__":
    create_guide_docx(r"c:\Users\swann\Desktop\Proejt gravity\Site 2026-2027\GUIDE_UTILISATION_PORTAIL.docx")
